"""Utility functions for reading Excel spreadsheets."""

from pathlib import Path
from typing import List, Dict, Any

try:
    from openpyxl import load_workbook
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False


def read_excel_to_dicts(file_path: str, sheet_name: int = 0) -> List[Dict[str, Any]]:
    """
    Read an Excel spreadsheet and return a list of dictionaries.

    The first row of the spreadsheet is used as field names.

    Args:
        file_path: Path to the Excel file (.xlsx)
        sheet_name: Index of the sheet to read (default: 0 for first sheet)

    Returns:
        List of dictionaries where keys are column names from the first row
        and values are the cell values from subsequent rows.

    Example:
        >>> data = read_excel_to_dicts('football_analytics/data/team_stats_definition.xlsx')
        >>> print(data[0])
        {
            'name': 'team_season_aggression',
            'type': 'double',
            'use': 'No',
            'group': 'Off-Ball',
            ...
        }
    """
    if not HAS_OPENPYXL:
        raise ImportError("openpyxl is required. Install with: pip install openpyxl")

    file_path = Path(file_path)

    if not file_path.exists():
        raise FileNotFoundError(f"Excel file not found: {file_path}")

    # Load the workbook
    wb = load_workbook(str(file_path))

    # Get the specified sheet
    if isinstance(sheet_name, int):
        ws = wb.worksheets[sheet_name]
    else:
        ws = wb[sheet_name]

    # Read all rows
    rows = list(ws.iter_rows(values_only=True))

    if not rows:
        return []

    # First row contains field names
    field_names = rows[0]

    # Convert to list of dictionaries
    result = []
    for row in rows[1:]:
        # Skip empty rows
        if not any(cell is not None for cell in row):
            continue

        # Create dictionary with field names as keys
        row_dict = {}
        for i, field_name in enumerate(field_names):
            if field_name is not None:
                value = row[i] if i < len(row) else None
                row_dict[field_name] = value

        result.append(row_dict)

    return result


def read_excel_sheet_names(file_path: str) -> List[str]:
    """
    Get all sheet names from an Excel file.

    Args:
        file_path: Path to the Excel file (.xlsx)

    Returns:
        List of sheet names
    """
    if not HAS_OPENPYXL:
        raise ImportError("openpyxl is required. Install with: pip install openpyxl")

    file_path = Path(file_path)

    if not file_path.exists():
        raise FileNotFoundError(f"Excel file not found: {file_path}")

    wb = load_workbook(str(file_path))
    return wb.sheetnames
